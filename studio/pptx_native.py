"""Independent deterministic PPTX builder (python-pptx, no npm/API required).

The design reference is NEVER silently inserted in a supposed fully-editable deck.
Only an approved or explicitly supplied background is used, while editable text
and shapes remain proper OOXML objects. Complex decorative art remains raster.
"""
from __future__ import annotations
from copy import deepcopy
import re
from pathlib import Path
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from .core import FILENAME_RE, find_slide, project_dir
from .fonts import default_cjk_font

W, H = 13.333333, 7.5


def east_asian_font() -> str:
    """Return a CJK family available to the local Office/LibreOffice renderer."""
    return default_cjk_font()


def _set_run_fonts(run, latin: str, east_asian: str) -> None:
    # python-pptx's font.name only writes a:latin. Office and LibreOffice use
    # a:ea for CJK glyphs, otherwise Chinese silently falls back to a missing font.
    run.font.name = latin
    rpr = run._r.get_or_add_rPr()
    for tag, family in (('a:latin', latin), ('a:ea', east_asian), ('a:cs', east_asian)):
        node = rpr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            rpr.append(node)
        node.set('typeface', family)

_MEASURE_FONT_FILES={
    'Noto Sans':'NotoSans-VF.ttf','Noto Serif':'NotoSerif-VF.ttf',
    'Noto Sans SC':'NotoSansSC-VF.ttf','Noto Serif SC':'NotoSerifSC-VF.ttf',
}
_CJK_FONT_FAMILIES={
    'Noto Sans SC','Noto Serif SC','Microsoft YaHei','SimSun','SimHei','KaiTi',
    'FangSong','Songti SC','HarmonyOS Sans SC','Hiragino Sans GB','PingFang SC',
    'Arial Unicode MS',
}

def _east_asian_family(style:dict,project:dict)->str:
    """Keep CJK glyphs in the same selected family when it supports Chinese."""
    explicit=style.get('eastAsianFont')
    if explicit:return str(explicit)
    latin=style.get('font') or project.get('design',{}).get('font')
    return str(latin) if latin in _CJK_FONT_FAMILIES else east_asian_font()

def _estimated_line_width_pt(text: str, size: float, tracking: float,
                             font_name: str|None=None,east_asian_font_name: str|None=None) -> float:
    """Measure bundled Noto glyph advances, falling back to a conservative estimate."""
    font_dir=Path(__file__).parent/'assets'/'fonts'
    latin_path=font_dir/_MEASURE_FONT_FILES.get(font_name or '', '')
    east_path=font_dir/_MEASURE_FONT_FILES.get(east_asian_font_name or '', '')
    if latin_path.is_file() and east_path.is_file():
        try:
            pixels=max(1,round(float(size)*96/72))
            latin=ImageFont.truetype(str(latin_path),pixels)
            east=ImageFont.truetype(str(east_path),pixels)
            runs=[];last_font=None;buffer=''
            for char in str(text):
                font=east if ord(char)>255 else latin
                if font is not last_font and buffer:
                    runs.append((last_font,buffer));buffer=''
                last_font=font;buffer+=char
            if buffer:runs.append((last_font,buffer))
            width_px=sum(font.getlength(run) for font,run in runs)
            return width_px*.75+max(0,len(text)-1)*float(tracking)
        except (OSError,ValueError,TypeError):
            pass
    ems=sum(1.0 if ord(char)>255 else (.34 if char.isspace() else .58) for char in text)
    return ems*size+max(0,len(text)-1)*tracking


def _measured_script_advance_pt(text: str, size: float, tracking: float,
                                font_name: str|None=None,east_asian_font_name: str|None=None) -> float:
    """Measure positioning advances with bundled Latin/CJK proxy faces.

    This is used only to position adjacent OCR script runs. The conservative
    estimator above remains responsible for deciding whether a line fits.
    """
    font_dir=Path(__file__).parent/'assets'/'fonts'
    latin_path=font_dir/_MEASURE_FONT_FILES.get(font_name or '', 'NotoSans-VF.ttf')
    east_path=font_dir/_MEASURE_FONT_FILES.get(east_asian_font_name or '', 'NotoSansSC-VF.ttf')
    if not latin_path.is_file():latin_path=font_dir/'NotoSans-VF.ttf'
    if not east_path.is_file():east_path=font_dir/'NotoSansSC-VF.ttf'
    try:
        pixels=max(1,round(float(size)*96/72))
        latin=ImageFont.truetype(str(latin_path),pixels)
        east=ImageFont.truetype(str(east_path),pixels)
        runs=[];last_font=None;buffer=''
        for char in str(text):
            font=east if ord(char)>255 else latin
            if font is not last_font and buffer:
                runs.append((last_font,buffer));buffer=''
            last_font=font;buffer+=char
        if buffer:runs.append((last_font,buffer))
        return sum(font.getlength(run) for font,run in runs)*.75+max(0,len(text)-1)*float(tracking)
    except (OSError,ValueError,TypeError):
        return _estimated_line_width_pt(text,size,tracking,font_name,east_asian_font_name)

def _single_line_texts(text: str, width_pt: float, size: float, tracking: float,
                       font_name: str|None=None,east_asian_font_name: str|None=None) -> list[str]:
    """Keep one visual line intact; hard line breaks become separate shapes.

    Automatically wrapping a word into an extra shape is unsafe here: OCR
    boxes are often only a few points narrower than their text, and dividing
    the original box height among those artificial fragments makes the final
    line boxes shorter than the font itself. The caller fits the font to the
    available width instead.
    """
    paragraphs=re.split(r'\r\n|[\r\n\v\f\u0085\u2028\u2029]',str(text or ''))
    return paragraphs or ['']


def _fit_single_line_font_size(text: str, width_pt: float, size: float, tracking: float,
                               font_name: str|None=None,
                               east_asian_font_name: str|None=None) -> float:
    """Conservatively shrink a single line so Office cannot wrap or clip it."""
    capacity=max(1.0,float(width_pt))
    required=_estimated_line_width_pt(text,size,tracking,font_name,east_asian_font_name)
    if required<=capacity:return float(size)
    # Leave room for Office/LibreOffice font substitution and text-frame
    # rounding, even when our estimated width is just below the box width.
    return max(4.0,min(float(size),float(size)*capacity*.80/max(required,1e-6)))


def _reflow_spatial_script_rows(project: dict) -> dict:
    """Close OCR-created gaps between adjacent script runs on the same line."""
    result=deepcopy(project)
    for page in result.get('slides',[]):
        buckets={}
        for obj in page.get('elements',[]):
            if obj.get('kind')!='text' or not obj.get('source_text'):
                continue
            box=obj.get('box') or []
            if len(box)!=4:continue
            key=re.sub(r'\s+','',str(obj['source_text'])).casefold()
            buckets.setdefault(key,[]).append(obj)
        page_rows={}
        for obj in page.get('elements',[]):
            group=obj.get('source_group')
            if obj.get('kind')=='text' and isinstance(group,str) and group:
                page_rows.setdefault(group,[]).append(obj)
        for key,items in buckets.items():
            items.sort(key=lambda obj:(float(obj['box'][1]),float(obj['box'][0])))
            clusters=[]
            for obj in items:
                x,y,w,h=map(float,obj['box']);height=max(.001,h);center=y+h/2
                cluster=next((c for c in clusters
                    if abs(c['center']-center)<=max(.008,min(c['height'],height)*.35)
                    and x-c['right']<=max(.018,min(c['height'],height)*1.4)),None)
                if cluster is None:
                    cluster={'items':[],'left':x,'right':x+w,'top':y,'bottom':y+h,'center':center,'height':height}
                    clusters.append(cluster)
                cluster['items'].append(obj);cluster['left']=min(cluster['left'],x)
                cluster['right']=max(cluster['right'],x+w);cluster['top']=min(cluster['top'],y)
                cluster['bottom']=max(cluster['bottom'],y+h);cluster['height']=cluster['bottom']-cluster['top']
                cluster['center']=(cluster['top']+cluster['bottom'])/2
            for index,cluster in enumerate(clusters):
                runs=cluster['items']
                if len(runs)<2:continue
                runs.sort(key=lambda obj:float(obj['box'][0]))
                if len({obj.get('text_class') for obj in runs})<2:continue
                size=min(float((obj.get('style') or {}).get('size',18)) for obj in runs)
                # OCR often gives a mixed-script bullet a shorter crop and a
                # smaller font size than neighboring bullets in the same
                # column. Borrow the column's measured line width and common
                # size from nearby full-length rows, so the mixed row does not
                # end up visibly smaller just because it has separate runs.
                siblings=[]
                own_groups={o.get('source_group') for o in runs}
                for group_id,other_items in page_rows.items():
                    if group_id in own_groups:continue
                    source=other_items[0].get('source_box') or other_items[0].get('box') or []
                    if len(source)!=4:continue
                    ox,oy,ow,oh=map(float,source);ocenter=oy+oh/2
                    other_sizes=[float((o.get('style') or {}).get('size',18)) for o in other_items]
                    if abs(ox-cluster['left'])<=.025 and abs(ocenter-cluster['center'])<=.16 \
                        and .7<=oh/max(cluster['height'],.001)<=1.4 \
                        and max(other_sizes,default=99)<=18 \
                        and max((len(str(o.get('source_text',''))) for o in other_items),default=0)>=12:
                        siblings.append({'items':other_items,'left':ox,'right':ox+ow,
                            'center':ocenter,'height':oh})
                if siblings:
                    peer_sizes=[float((o.get('style') or {}).get('size',size))
                        for other in siblings for o in other['items']
                        if 8<=float((o.get('style') or {}).get('size',size))<=32]
                    if peer_sizes:size=max(size,sorted(peer_sizes)[len(peer_sizes)//2])
                    if cluster['left']<.5:
                        next_column=min((float((items[0].get('source_box') or items[0].get('box'))[0])
                            for items in page_rows.values()
                            if len(items[0].get('source_box') or items[0].get('box') or [])==4
                            and float((items[0].get('source_box') or items[0].get('box'))[0])>cluster['left']+.12
                            and abs(float((items[0].get('source_box') or items[0].get('box'))[1])+
                                float((items[0].get('source_box') or items[0].get('box'))[3])/2-cluster['center'])<=.18
                            and max((len(str(o.get('source_text',''))) for o in items),default=0)>=12),default=1.0)
                        available_right=min(.96,next_column-.035)
                    else:
                        available_right=.96
                    cluster_right=max(cluster['right'],*(other['right'] for other in siblings),available_right)
                else:
                    cluster_right=cluster['right']
                advances=[]
                for obj in runs:
                    style=obj.get('style') or {}
                    advance=_estimated_line_width_pt(str(obj.get('text','')),size,
                        float(style.get('char_spacing_pt',0)),style.get('font') or result.get('design',{}).get('font'),
                        _east_asian_family(style,result))
                    advances.append(advance)
                # Script-specific textboxes remain separate, but their glyphs
                # are laid out with zero artificial inter-script spacing.
                gaps=[0.0 for _ in range(len(runs)-1)]
                capacity=(cluster_right-cluster['left'])*W*72
                total=sum(advances)+sum(gaps)
                # The source OCR boxes already describe the line's available
                # width. A second 18% safety shrink here made mixed-script
                # text visibly smaller than neighboring text. Fit only when
                # measured advances genuinely exceed the source line width.
                fitted=max(4.0,size*min(1.0,capacity/max(total,1e-6)))
                union=[cluster['left'],cluster['top'],cluster_right-cluster['left'],cluster['bottom']-cluster['top']]
                group_id=f"ocr-spatial:{key[:72]}:{round(cluster['center']*10000)}:{index}"
                cursor=cluster['left']
                for i,obj in enumerate(runs):
                    style=obj.setdefault('style',{})
                    latin=style.get('font') or result.get('design',{}).get('font')
                    east=_east_asian_family(style,result)
                    conservative=_estimated_line_width_pt(str(obj.get('text','')),fitted,
                        float(style.get('char_spacing_pt',0)),latin,east)
                    advance=_measured_script_advance_pt(str(obj.get('text','')),fitted,
                        float(style.get('char_spacing_pt',0)),latin,east)/(W*72)
                    if obj.get('text_class')=='latin':
                        # A split Latin run loses the in-line font's boundary
                        # kerning context; Noto's proxy advance is slightly
                        # wider than Office's YaHei substitution. Pull the
                        # following run in by 8% while keeping textbox slack.
                        advance*=.92
                    width=max(advance*1.10,conservative*1.08/(W*72))  # invisible box slack; next glyph starts at measured advance
                    obj['box']=[round(cursor,7),round(cluster['bottom']-union[3],7),round(width,7),round(union[3],7)]
                    obj['source_box']=list(union);obj['source_group']=group_id
                    style['size']=round(fitted,1);style['vertical_align']='bottom'
                    cursor+=advance+(gaps[i]/(W*72) if i<len(gaps) else 0)
    return result


def _fit_grouped_script_runs(project: dict) -> dict:
    """Apply one conservative point size to every script box in an OCR row."""
    result=deepcopy(project)
    for page in result.get('slides',[]):
        groups={}
        for obj in page.get('elements',[]):
            group=obj.get('source_group')
            if obj.get('kind')=='text' and isinstance(group,str) and group:
                groups.setdefault(group,[]).append(obj)
        for group in groups.values():
            if len(group)<2:continue
            current=min(float(obj.get('style',{}).get('size',18)) for obj in group)
            total_width=0.0
            available_width=None
            for obj in group:
                style=obj.get('style') or {};box=obj['box']
                source_box=obj.get('source_box')
                if isinstance(source_box,list) and len(source_box)==4:
                    available_width=float(source_box[2])*W*72
                width=_estimated_line_width_pt(str(obj.get('text','')),current,
                    float(style.get('char_spacing_pt',0)),style.get('font') or result.get('design',{}).get('font'),
                    _east_asian_family(style,result))
                total_width+=width
            fitted=max(4.0,min(current,current*(available_width or total_width)/max(total_width,1e-6)))
            if fitted<current:
                for obj in group:obj.setdefault('style',{})['size']=round(fitted,1)
    return result


def color(x, fallback='000000'):
    v=str(x or fallback).lstrip('#')
    if not re.fullmatch('[0-9a-fA-F]{6}',v):v=fallback
    return RGBColor.from_string(v.upper())

def coordinates(box):
    if len(box)!=4 or any(not isinstance(x,(int,float)) for x in box):
        raise ValueError('非法元素坐标')
    x,y,w,h=map(float,box)
    if min(x,y)<0 or min(w,h)<=0 or x+w>1.0001 or y+h>1.0001:
        raise ValueError('元素超出页面或尺寸不正确')
    return Inches(W*x),Inches(H*y),Inches(W*w),Inches(H*h)

def asset_path(project, sid, filename):
    if not FILENAME_RE.fullmatch(filename or '') or not filename.upper().startswith(sid+'_'):
        raise ValueError('跨页或非法媒体引用：'+str(filename))
    path=project_dir(project['id'])/'images'/filename
    if not path.is_file():raise FileNotFoundError('图片不存在：'+filename)
    return path

def add_image(slide, path, box, fit='stretch'):
    x,y,w,h=coordinates(box)
    if fit not in ('contain','cover','stretch'):fit='stretch'
    if fit=='stretch':return slide.shapes.add_picture(str(path),x,y,width=w,height=h)
    with Image.open(path) as im: iw,ih=im.size
    if fit=='contain':
        scale=min(w/iw,h/ih);dw=int(iw*scale);dh=int(ih*scale)
        return slide.shapes.add_picture(str(path),x+(w-dw)//2,y+(h-dh)//2,width=dw,height=dh)
    scale=max(w/iw,h/ih);dw=int(iw*scale);dh=int(ih*scale)
    shape=slide.shapes.add_picture(str(path),x,y,width=w,height=h)
    if dw>w: shape.crop_left=shape.crop_right=(1-w/dw)/2
    if dh>h: shape.crop_top=shape.crop_bottom=(1-h/dh)/2
    return shape

def build(project:dict, output:Path, mode='editable'):
    if mode not in ('editable','raster'):raise ValueError('非法导出模式')
    # _reflow_spatial_script_rows already applies the single necessary width
    # fit. Running a second grouped fit here shrank every OCR row twice.
    project=_reflow_spatial_script_rows(project) if mode=='editable' else project
    ppt=Presentation();ppt.slide_width=Inches(W);ppt.slide_height=Inches(H)
    for page in project['slides']:
        sid=page['id'];s=ppt.slides.add_slide(ppt.slide_layouts[6])
        bg=page.get('background') or project['design']['bg']
        s.background.fill.solid();s.background.fill.fore_color.rgb=color(bg,'FFFFFF')
        imgs=page.get('images',{})
        if mode=='raster':
            if imgs.get('design'): add_image(s,asset_path(project,sid,imgs['design']),[0,0,1,1])
            continue
        if imgs.get('clean'):
            meta=page.get('image_metadata',{}).get('clean',{})
            if meta.get('stale_source'):raise ValueError(f'{sid} 无字背景与新版设计稿不对应，请重新审核')
            add_image(s,asset_path(project,sid,imgs['clean']),[0,0,1,1])
        for obj in page.get('elements',[]):
            kind=obj.get('kind'); st=obj.get('style') or {}
            x,y,w,h=coordinates(obj['box'])
            if kind=='text':
                size=max(4,min(160,float(st.get('size',18))))
                render_text=str(obj.get('text',''))
                if obj.get('text_class')=='latin':
                    # Keep full-width CJK punctuation from reintroducing an
                    # East-Asian fallback font inside an otherwise Latin-only
                    # box. NFKC-equivalent punctuation preserves the text
                    # check while keeping the Latin glyphs on one baseline.
                    render_text=render_text.translate(str.maketrans({'（':'(','）':')','，':',','：':':','；':';','？':'?','！':'!'}))
                east_font=_east_asian_family(st,project)
                tracking=max(-5,min(20,float(st.get('char_spacing_pt',0))))
                margin_left=max(0,min(72,float(st.get('margin_left_pt',0))))
                margin_right=max(0,min(72,float(st.get('margin_right_pt',0))))
                line_texts=_single_line_texts(render_text,float(obj['box'][2])*W*72-margin_left-margin_right,size,tracking,
                                              st.get('font') or project['design'].get('font'),
                                              east_font)
                available_width=float(obj['box'][2])*W*72-margin_left-margin_right
                fitted_size=min((_fit_single_line_font_size(line,available_width,size,tracking,
                                                            st.get('font') or project['design'].get('font'),east_font)
                                 for line in line_texts if line.strip()),default=size)
                slot_h=h/max(1,len(line_texts))
                for i,line in enumerate(line_texts):
                    if not line.strip():continue
                    line_y=y+int(i*slot_h)
                    line_height=int(slot_h) if len(line_texts)>1 else h
                    tb=s.shapes.add_textbox(x,line_y,w,line_height)
                    base_name=str(obj.get('id','text'))
                    tb.name=base_name if len(line_texts)==1 else f'{base_name}_line_{i+1:02d}'
                    tf=tb.text_frame;tf.clear();tf.word_wrap=False
                    tf.margin_left=Pt(margin_left)
                    tf.margin_right=Pt(margin_right)
                    tf.margin_top=Pt(max(0,min(72,float(st.get('margin_top_pt',0)))))
                    tf.margin_bottom=Pt(max(0,min(72,float(st.get('margin_bottom_pt',0)))))
                    tf.vertical_anchor={'top':MSO_ANCHOR.TOP,'middle':MSO_ANCHOR.MIDDLE,'bottom':MSO_ANCHOR.BOTTOM}.get(st.get('vertical_align'),MSO_ANCHOR.TOP)
                    p=tf.paragraphs[0];p.text=line
                    p.line_spacing=max(.7,min(2.5,float(st.get('line_spacing',1.05))))
                    p.space_before=Pt(max(0,min(72,float(st.get('paragraph_space_before_pt',0)))))
                    p.space_after=Pt(max(0,min(72,float(st.get('paragraph_space_after_pt',0)))))
                    p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(st.get('align'),PP_ALIGN.LEFT)
                    for r in p.runs:
                        latin_font=st.get('font') or project['design'].get('font') or 'Arial'
                        _set_run_fonts(r,latin_font,east_font)
                        r.font.size=Pt(fitted_size)
                        r.font.bold=bool(st.get('bold',False))
                        r.font.italic=bool(st.get('italic',False))
                        r.font.color.rgb=color(st.get('color'),project['design']['ink'])
                        # DrawingML run spacing is in 1/100 pt; python-pptx
                        # currently has no public Font.spacing property.
                        if abs(tracking)>1e-6:
                            r._r.get_or_add_rPr().set('spc',str(round(tracking*100)))
            elif kind=='formula':
                from .equations import append_omml
                tb=s.shapes.add_textbox(x,y,w,h);tf=tb.text_frame
                tf.clear();tf.word_wrap=True
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                tf.vertical_anchor=MSO_ANCHOR.MIDDLE
                paragraph=tf.paragraphs[0]
                append_omml(paragraph,str(obj.get('text','')),float(st.get('size',22)),st.get('align','center'))
            elif kind in ('rect','ellipse'):
                preset=MSO_SHAPE.OVAL if kind=='ellipse' else MSO_SHAPE.RECTANGLE
                sh=s.shapes.add_shape(preset,x,y,w,h);sh.fill.solid()
                sh.fill.fore_color.rgb=color(st.get('fill'),project['design']['accent'])
                sh.line.color.rgb=color(st.get('stroke'),st.get('fill') or project['design']['accent'])
                sh.line.width=Pt(max(0,float(st.get('strokeWidth',0))))
            elif kind=='line':
                from pptx.enum.shapes import MSO_CONNECTOR
                sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,x,y,x+w,y+h)
                sh.line.color.rgb=color(st.get('stroke'),project['design']['accent'])
                sh.line.width=Pt(max(.25,float(st.get('strokeWidth',2))))
                if st.get('arrow'):
                    from pptx.oxml.xmlchemy import OxmlElement
                    line_props=sh._element.spPr.get_or_add_ln()
                    end=OxmlElement('a:tailEnd');end.set('type','triangle');line_props.append(end)
            elif kind=='image':
                add_image(s,asset_path(project,sid,obj['asset']),obj['box'],st.get('fit','stretch'))
            else:raise ValueError('不支持的对象类型：'+str(kind))
    output.parent.mkdir(parents=True,exist_ok=True);ppt.save(str(output))
    return {'slides':len(ppt.slides),'backend':'python-pptx','native_elements':sum(len(s.shapes) for s in ppt.slides)}
