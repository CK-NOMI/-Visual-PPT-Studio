"""Deterministic QA. A score is not manufactured from visual similarity."""
from pathlib import Path
import math
import re
import unicodedata
from PIL import Image, ImageStat
import numpy as np
from .core import project_dir


def intersects(a,b):
    x1=max(a[0],b[0]);y1=max(a[1],b[1]);x2=min(a[0]+a[2],b[0]+b[2]);y2=min(a[1]+a[3],b[1]+b[3])
    if x2<=x1 or y2<=y1:return 0
    return (x2-x1)*(y2-y1)


def visible_glyph_box(element):
    """Predict ink bounds, not transparent frame capacity; render checks remain separate."""
    layout=element.get('text_layout') or {}
    if layout.get('version')==2 and layout.get('ink_box_pt'):
        x,y,w,h=element['box'];st=element.get('style') or {};ink=layout['ink_box_pt']
        baseline=layout['baseline_pt']
        left=float(st.get('margin_left_pt',0));right=float(st.get('margin_right_pt',0))
        offset=(w*960-left-right-layout.get('advance_pt',0))*{'left':0,'center':.5,'right':1}.get(st.get('align','left'),0)
        return [(x*960+left+offset+ink[0])/960,
                (baseline+ink[1])/540,max(.00001,ink[2]-ink[0])/960,
                max(.00001,ink[3]-ink[1])/540]
    return element['box']

def _luminance(rgb):
    vals=[]
    for v in rgb:
        x=v/255.0
        vals.append(x/12.92 if x<=0.04045 else ((x+0.055)/1.055)**2.4)
    return .2126*vals[0]+.7152*vals[1]+.0722*vals[2]

def _contrast(hx,rgb):
    try: fg=tuple(int(str(hx).strip('#')[j:j+2],16) for j in (0,2,4))
    except Exception:return 21.0
    a,b=sorted([_luminance(fg),_luminance(rgb)])
    return (b+.05)/(a+.05)

def _estimated_lines(text,width_pt,font_pt,char_spacing_pt=0):
    if not text:return 0
    # Estimate glyph advances in ems, then include PowerPoint tracking (points)
    # between characters. It is intentionally conservative for mixed scripts.
    tracking_em=float(char_spacing_pt or 0)/max(font_pt,1)
    total=0
    for line in text.split('\n'):
        # Approximate; non-Latin glyphs are close to 1em, Latin letters often narrower.
        advances=[1.0 if ord(c)>255 else (.4 if c.isspace() else .58) for c in line]
        units=sum(advances)+max(0,len(line)-1)*tracking_em
        capacity=max(width_pt/max(font_pt,1),1)
        total+=max(1,math.ceil(units/capacity))
    return total

def audit(project):
    warnings=[]; errors=[]; slides_report=[]
    for s in project['slides']:
        sid=s['id']; issues=[]; els=s.get('elements',[])
        if 'design' not in s['images']:issues.append(('warning','缺少完整设计稿'))
        if 'clean' not in s['images']:issues.append(('warning','缺少无文字背景：混合导出将使用纯色背景，不会保留设计图'))
        clean_meta=s.get('image_metadata',{}).get('clean',{})
        if clean_meta.get('stale_source'):issues.append(('error','设计图替换后底图已过期：需重建/重新上传或明确重新审核'))
        if clean_meta.get('embedded_text_remaining'):issues.append(('warning','选择性混合底图仍包含栅格文字：仅确认提升的文字可编辑，不能声称全页文字可编辑'))
        if not any(e['kind']=='text' for e in els):issues.append(('warning','没有可编辑的文本框'))
        for role,fname in s['images'].items():
            p=project_dir(project['id'])/'images'/fname
            if not p.is_file():issues.append(('error',f'{role} 引用的文件不存在'))
            else:
                try:
                    with Image.open(p) as im:
                        r=im.width/im.height
                        if role in ('design','clean') and abs(r-16/9)>.06:issues.append(('warning',f'{role} 图片不是16:9，将拉伸到16:9'))
                except Exception:issues.append(('error',f'{role} 图片无法打开'))
        clean=s['images'].get('clean')
        clean_img=None
        if clean:
            try:clean_img=Image.open(project_dir(project['id'])/'images'/clean).convert('RGB')
            except Exception:pass
        if s['images'].get('design') and s['images'].get('clean'):
            try:
                img_dir=project_dir(project['id'])/'images'
                with Image.open(img_dir/s['images']['design']) as ref, Image.open(img_dir/s['images']['clean']) as cleaned:
                    a=np.asarray(ref.convert('RGB').resize((256,144)),dtype=np.int16)
                    b=np.asarray(cleaned.convert('RGB').resize((256,144)),dtype=np.int16)
                region=np.ones((144,256),dtype=bool)
                for e in els:
                    if e['kind']=='text':
                        ex,ey,ew,eh=visible_glyph_box(e)
                        # Exclude text areas: two images SHOULD differ here.
                        x0=max(0,int((ex-.02)*256));y0=max(0,int((ey-.02)*144))
                        x1=min(256,int((ex+ew+.02)*256));y1=min(144,int((ey+eh+.02)*144))
                        region[y0:y1,x0:x1]=False
                if region.any():
                    diff=float(np.mean(np.abs(a[region]-b[region]))/255)
                    if diff>.22:
                        issues.append(('warning',f'设计稿与去字底图在非文字区域差异明显（RGB平均差异{diff:.2f}）；可能不适合直接重建'))
            except Exception:
                issues.append(('warning','无法评估设计稿和无文字背景的一致性'))
        for e in els:
            x,y,w,h=e.get('box',[0,0,0,0])
            # Off-canvas native objects are legal in PowerPoint (for clipping,
            # bleed, and partially hidden elements). Report them for review,
            # but do not make an otherwise valid optimized deck impossible to
            # export solely because an object extends past the slide edge.
            if x<0 or y<0 or x+w>1.001 or y+h>1.001:issues.append(('warning',f'{e["id"]} 部分超出页面边界；PowerPoint 会按页面裁切显示'))
            if e['kind']=='text' and not e.get('text','').strip():issues.append(('warning',f'{e["id"]} 文本为空'))
            if e['kind']=='text' and w>0 and h>0:
                st=e.get('style',{})
                fontsize=float(st.get('size',18))
                available_w=w*13.333*72-float(st.get('margin_left_pt',0))-float(st.get('margin_right_pt',0))
                available_h=h*7.5*72-float(st.get('margin_top_pt',0))-float(st.get('margin_bottom_pt',0))
                lines=_estimated_lines(e.get('text',''),available_w,fontsize,st.get('char_spacing_pt',0))
                paragraph_gaps=max(0,str(e.get('text','')).count('\n'))*float(st.get('paragraph_space_after_pt',0))
                if e.get('text_layout',{}).get('version')==2:
                    for issue in e['text_layout'].get('layout_warning',[]):
                        issues.append(('warning',f'{e["id"]} 布局诊断：{issue}'))
                elif lines*fontsize*1.22*float(st.get('line_spacing',1.05))+paragraph_gaps>available_h*1.15:
                    issues.append(('warning',f'{e["id"]} 文本可能溢出：建议扩大文本框或缩小字号'))
                if clean_img and all(v>=0 for v in (x,y,w,h)) and x+w<=1.0 and y+h<=1.0:
                    # Sample actual cleaned background under the text area.
                    x,y,w,h=visible_glyph_box(e)
                    imw,imh=clean_img.size
                    sample=clean_img.crop((int(x*imw),int(y*imh),max(int((x+w)*imw),int(x*imw)+1),max(int((y+h)*imh),int(y*imh)+1)))
                    sample.thumbnail((45,35))
                    brightness=tuple(int(c) for c in ImageStat.Stat(sample).median[:3])
                    ratio=_contrast(st.get('color',project['design']['ink']),brightness)
                    if ratio<3.0:issues.append(('warning',f'{e["id"]} 文字与背景对比度可能偏低（粗略比值{ratio:.1f}）'))
            if e['kind']=='image' and not (project_dir(project['id'])/'images'/e.get('asset','')).is_file():issues.append(('error',f'{e["id"]} 素材丢失'))
        # warn substantial overlap of text objects only, not decorative regions
        texts=[e for e in els if e['kind']=='text']
        for i,a in enumerate(texts):
            for b in texts[i+1:]:
                ab,bb=visible_glyph_box(a),visible_glyph_box(b)
                overlap=intersects(ab,bb)
                if overlap>0.03*min(ab[2]*ab[3],bb[2]*bb[3]):
                    issues.append(('warning',f'{a["id"]} 与 {b["id"]} 的可见文字可能重叠'))
        for sev,msg in issues:
            (errors if sev=='error' else warnings).append({'slide_id':sid,'message':msg})
        slides_report.append({'slide_id':sid,'errors':[m for k,m in issues if k=='error'],
                              'warnings':[m for k,m in issues if k=='warning']})
    return {'ok':len(errors)==0,'errors':errors,'warnings':warnings,'pages':slides_report,
            'counts':{'errors':len(errors),'warnings':len(warnings)},
            'meaning':'几何、文件和资源规则审计；不代表审美评分或文字溢出已完全排除。'}

def check_native_text(project,pptx_path):
    from pptx import Presentation
    import zipfile
    pp=Presentation(pptx_path);reports=[]
    if len(pp.slides)!=len(project['slides']):raise ValueError('导出页数不正确')
    with zipfile.ZipFile(pptx_path) as archive:
        from lxml import etree
        math_ns='{http://schemas.openxmlformats.org/officeDocument/2006/math}'
        slide_xml={f'ppt/slides/slide{i}.xml':etree.fromstring(archive.read(f'ppt/slides/slide{i}.xml'))
                   for i in range(1,len(pp.slides)+1)}
    for index,(expected,real) in enumerate(zip(project['slides'],pp.slides),1):
        actual=[sh for sh in real.shapes if sh.has_text_frame]
        actual_by_id={}
        for shape in actual:
            name=str(shape.name or '')
            match=re.fullmatch(r'(.+)_line_(\d+)',name)
            object_id=match.group(1) if match else name
            actual_by_id.setdefault(object_id,[]).append(shape.text)
        def text_key(value):
            # Office may expose line-split text boxes separately and may
            # normalize ordinary spaces. Compare each logical object after
            # joining its generated line shapes and removing whitespace.
            return ''.join(unicodedata.normalize('NFKC',str(value or '')).casefold().split())
        needed=[e for e in expected['elements'] if e['kind']=='text' and e.get('text')]
        missing=[]
        for element in needed:
            parts=actual_by_id.get(str(element['id']),[])
            observed=''.join(parts) if parts else ''
            if text_key(observed)!=text_key(element['text']):
                missing.append(element['text'])
        equations=[e for e in expected['elements'] if e['kind']=='formula' and e.get('text')]
        equation_count=sum(1 for node in slide_xml[f'ppt/slides/slide{index}.xml'].iter()
                           if node.tag==math_ns+'oMathPara')
        formula_pass=equation_count>=len(equations)
        reports.append({'slide_id':expected['id'],'native_text_count':len(actual),'expected_text_count':len(needed),
                        'missing':missing,'native_equation_count':equation_count,
                        'expected_equation_count':len(equations),'equations_pass':formula_pass,
                        'pass':not missing and formula_pass})
    return {'ok':all(x['pass'] for x in reports),'slides':reports}


def check_saved_layout(project,pptx_path):
    """Exact export contract, separate from loose OCR semantic matching."""
    from pptx import Presentation
    from pptx.oxml.ns import qn
    pp=Presentation(pptx_path);pages=[]
    for source,slide in zip(project['slides'],pp.slides):
        by_name={}
        for sh in slide.shapes:
            if sh.has_text_frame:by_name.setdefault(sh.name,[]).append(sh)
        expected={e['id']:e for e in source['elements'] if e['kind'] in ('text','formula') and str(e.get('text','')).strip()}
        differences=[]
        for name,shapes in by_name.items():
            if name not in expected and any(s.text.strip() for s in shapes):
                differences.append({'id':name,'type':'unexpected_object'})
            if len(shapes)>1:differences.append({'id':name,'type':'duplicate_object','count':len(shapes)})
        for id_,e in expected.items():
            shapes=by_name.get(id_,[])
            if len(shapes)!=1:
                differences.append({'id':id_,'type':'missing_or_duplicate_object'});continue
            sh=shapes[0];tf=sh.text_frame
            if e['kind']=='text' and sh.text!=e['text']:
                differences.append({'id':id_,'type':'exact_text_or_whitespace_mismatch','expected':e['text'],'actual':sh.text})
            actual=[sh.left/pp.slide_width,sh.top/pp.slide_height,sh.width/pp.slide_width,sh.height/pp.slide_height]
            if any(abs(a-b)>1e-6 for a,b in zip(actual,e['box'])):
                differences.append({'id':id_,'type':'export_changed_box','expected':e['box'],'actual':actual})
            if tf.word_wrap is not False or tf._txBody.bodyPr.find(qn('a:noAutofit')) is None:
                differences.append({'id':id_,'type':'wrap_or_autofit_enabled'})
            if e['kind']=='text':
                st=e.get('style',{})
                latin=st.get('font') or project.get('design',{}).get('font') or 'Arial'
                east=st.get('eastAsianFont') or latin
                multiline='\n' in e.get('text','')
                if len(tf.paragraphs)!= (e.get('text','').count('\n')+1 if multiline else 1):
                    differences.append({'id':id_,'type':'export_created_multiple_paragraphs'})
                expected_runs=e.get('rich_runs') or [{'text':e.get('text',''),'style':{}}]
                expected_segments=[]
                for fragment in expected_runs:
                    for part in fragment.get('text','').split('\n'):
                        if part:expected_segments.append({'text':part,'style':fragment.get('style',{})})
                actual_runs=[run for p in tf.paragraphs for run in p.runs]
                if [run.text for run in actual_runs]!=[part['text'] for part in expected_segments]:
                    differences.append({'id':id_,'type':'rich_run_text_mismatch',
                        'expected':[part['text'] for part in expected_segments],
                        'actual':[run.text for run in actual_runs]})
                for run_index,run in enumerate(actual_runs):
                    seg=expected_segments[run_index] if run_index<len(expected_segments) else {'style':{}}
                    run_style={**st,**seg.get('style',{})}
                    run_latin=run_style.get('font') or latin
                    run_east=run_style.get('eastAsianFont') or run_style.get('font') or east
                    size=run.font.size.pt if run.font.size else None
                    if size is None or abs(size-float(run_style.get('size',18)))>.011:
                        differences.append({'id':id_,'type':'export_changed_font_size','actual':size,
                                            'expected':run_style.get('size',18)})
                    props=run._r.get_or_add_rPr()
                    for tag,family in [('a:latin',run_latin),('a:ea',run_east),('a:cs',run_east)]:
                        node=props.find(qn(tag))
                        if node is None or node.get('typeface')!=family:
                            differences.append({'id':id_,'type':'export_changed_font','slot':tag,'expected':family,
                                                'actual':node.get('typeface') if node is not None else None})
                    if bool(run.font.bold)!=bool(run_style.get('bold',False)) or bool(run.font.italic)!=bool(run_style.get('italic',False)):
                        differences.append({'id':id_,'type':'export_changed_font_style'})
                    spacing=props.get('spc')
                    expected_spacing=round(float(run_style.get('char_spacing_pt',0))*100)
                    if (spacing is None and expected_spacing!=0) or (spacing is not None and int(spacing)!=expected_spacing):
                        differences.append({'id':id_,'type':'export_changed_character_spacing',
                            'expected':expected_spacing,'actual':spacing})
                    if bool(run.font.underline)!=bool(run_style.get('underline',False)):
                        differences.append({'id':id_,'type':'export_changed_underline'})
                    strike=props.get('strike') in {'sngStrike','dblStrike'}
                    if strike!=bool(run_style.get('strike',False)):
                        differences.append({'id':id_,'type':'export_changed_strike'})
                for p in tf.paragraphs:
                    try:
                        if p.line_spacing is None or abs(float(p.line_spacing)-float(st.get('line_spacing',1.05)))>.011:
                            differences.append({'id':id_,'type':'export_changed_line_spacing'})
                    except (TypeError,ValueError):
                        differences.append({'id':id_,'type':'export_changed_line_spacing'})
                for key,margin in [('margin_left_pt',tf.margin_left),('margin_right_pt',tf.margin_right),
                                   ('margin_top_pt',tf.margin_top),('margin_bottom_pt',tf.margin_bottom)]:
                    if abs(margin.pt-float(st.get(key,0)))>.011:
                        differences.append({'id':id_,'type':'export_changed_margin','field':key})
        pages.append({'slide_id':source['id'],'pass':not differences,'differences':differences})
    return {'ok':len(pp.slides)==len(project['slides']) and all(p['pass'] for p in pages),
            'pages':pages,'meaning':'精确文本/空白、富文本段、对象数量、坐标、字号、字距、行距、边距及禁自动缩放导出契约；WPS视觉验收另行记录。'}
