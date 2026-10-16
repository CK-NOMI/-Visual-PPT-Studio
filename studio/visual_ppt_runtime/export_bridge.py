"""Non-destructive compatibility adapter for Visual PPT Studio's python-pptx builder.

Why: legacy pptx_native.build produces a single-style run and rejects newline.
This adapter makes a temporary sanitized copy, builds with *existing* code,
then applies rich runs/paragraph styles/rotation to native PPT objects.
It never touches the original project.json or replaces Studio's public exporter.

To integrate production exports, host must call build_with_extensions instead of
studio.pptx_native.build; using unpatched exports with rich_runs is forbidden.
"""
from __future__ import annotations
from copy import deepcopy
from pathlib import Path
from typing import Callable
import re
from .errors import ValidationError
from .scene import validate_scene, validate_rich_runs


def _apply_fonts(run, latin, east):
    from pptx.oxml.xmlchemy import OxmlElement
    if latin:run.font.name=latin
    if east:
        props=run._r.get_or_add_rPr()
        for tag in ('a:ea','a:cs'):
            node=props.find('{http://schemas.openxmlformats.org/drawingml/2006/main}'+tag.split(':')[1])
            if node is None:node=OxmlElement(tag);props.append(node)
            node.set('typeface',east)


def _apply_run(run, defaults, overrides):
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    style={**defaults,**overrides}
    _apply_fonts(run,style.get('font'),style.get('eastAsianFont') or style.get('font'))
    if style.get('size') is not None:run.font.size=Pt(float(style['size']))
    for k in ('bold','italic','underline'):
        if k in style:setattr(run.font,k,bool(style[k]))
    if 'strike' in style:
        run._r.get_or_add_rPr().set('strike','sngStrike' if style['strike'] else 'noStrike')
    if 'color' in style:
        c=str(style['color']).lstrip('#');run.font.color.rgb=RGBColor.from_string(c.upper())
    if 'char_spacing_pt' in style:
        run._r.get_or_add_rPr().set('spc',str(round(float(style['char_spacing_pt'])*100)))


def _new_paragraph_style(paragraph,style):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    paragraph.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(style.get('align','left'),PP_ALIGN.LEFT)
    paragraph.line_spacing=float(style.get('line_spacing',1.05))
    paragraph.space_before=Pt(float(style.get('paragraph_space_before_pt',0)))
    paragraph.space_after=Pt(float(style.get('paragraph_space_after_pt',0)))


def _rewrite_text(shape,obj):
    """One DrawingML text run per requested rich run, split on line feeds."""
    from pptx.enum.text import MSO_AUTO_SIZE,MSO_ANCHOR
    from pptx.util import Pt
    full=obj.get('text','')
    default=dict(obj.get('style') or {})
    runs=validate_rich_runs(obj.get('rich_runs'),full,optional=True)
    if runs is None:runs=[{'text':full,'style':{}}]
    tf=shape.text_frame
    tf.clear();tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=Pt(float(default.get('margin_left_pt',0)))
    tf.margin_right=Pt(float(default.get('margin_right_pt',0)))
    tf.margin_top=Pt(float(default.get('margin_top_pt',0)))
    tf.margin_bottom=Pt(float(default.get('margin_bottom_pt',0)))
    tf.vertical_anchor={'top':MSO_ANCHOR.TOP,'middle':MSO_ANCHOR.MIDDLE,'bottom':MSO_ANCHOR.BOTTOM}.get(default.get('vertical_align','top'),MSO_ANCHOR.TOP)
    paragraph=tf.paragraphs[0];_new_paragraph_style(paragraph,default)
    for fragment in runs:
        parts=fragment['text'].split('\n')
        for index,text in enumerate(parts):
            if index:
                paragraph=tf.add_paragraph();_new_paragraph_style(paragraph,default)
            if text:
                run=paragraph.add_run();run.text=text
                _apply_run(run,default,fragment.get('style',{}))
    assert '\n'.join(p.text for p in tf.paragraphs)==full


def patch_pptx_file(project:dict,pptx_path:Path):
    """Apply exporter extensions onto already-built PPTX. Atomic replacement."""
    from pptx import Presentation
    validate_scene(project)
    path=Path(pptx_path)
    if not path.is_file():raise ValidationError('PPTX file missing')
    deck=Presentation(str(path))
    if len(deck.slides)!=len(project['slides']):raise ValidationError('PPTX slide count mismatch')
    touched=[]
    for page,pslide in zip(project['slides'],deck.slides):
        shapes={sh.name:sh for sh in pslide.shapes}
        offset=1 if page.get('images',{}).get('clean') else 0
        ordinal_ready=len(pslide.shapes)==len(page['elements'])+offset
        for obj_index,obj in enumerate(page['elements']):
            kind=obj.get('kind');name=obj['id'];shape=shapes.get(name)
            patch_text=(kind=='text' and ('rich_runs' in obj or '\n' in obj.get('text','') or
                        any(k in obj.get('style',{}) for k in ('underline','strike'))))
            patch_rotation=('rotation' in obj)
            patch_crop=(kind=='image' and any(k.startswith('crop_') for k in (obj.get('style') or {})))
            if not (patch_text or patch_rotation or patch_crop):continue
            if shape is None and ordinal_ready:
                shape=pslide.shapes[obj_index+offset]
            if shape is None:raise ValidationError('PPTX shape missing for extension: '+name)
            if patch_text:
                if not shape.has_text_frame:raise ValidationError('native text shape missing: '+name)
                _rewrite_text(shape,obj)
            if patch_rotation:shape.rotation=float(obj['rotation'])
            if patch_crop:
                from pptx.enum.shapes import MSO_SHAPE_TYPE
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:raise ValidationError('image crop needs native picture shape: '+name)
                st=obj.get('style') or {}
                if st.get('crop_left',0)+st.get('crop_right',0)>=1 or st.get('crop_top',0)+st.get('crop_bottom',0)>=1:
                    raise ValidationError('image cropping leaves no picture area')
                for key in ('crop_left','crop_right','crop_top','crop_bottom'):
                    if key in st:setattr(shape,key,float(st[key]))
            touched.append({'slide_id':page['id'],'element_id':name,'rich_text':patch_text,'rotation':patch_rotation,'crop':patch_crop})
    staging=path.with_name('.'+path.name+'.patched')
    try:
        deck.save(str(staging))
        staging.replace(path)
    finally:staging.unlink(missing_ok=True)
    return {'patched_objects':touched,'count':len(touched)}


def build_with_extensions(project:dict,output:Path,mode='editable',*,base_builder:Callable|None=None):
    """Build using Studio legacy builder, then extend in-place.

    Any embedded legacy `text_layout` is invalidated for changed text/typography
    by EditingSession. For precision rendering, the PPTX is the visual truth;
    browser Stage 02 preview needs its own rich-run renderer integration.
    """
    validate_scene(project)
    if base_builder is None:
        from studio.pptx_native import build as base_builder
    prepared=deepcopy(project)
    if mode=='editable':
        for page in prepared['slides']:
            for obj in page['elements']:
                if obj.get('kind')=='text' and '\n' in obj.get('text',''):
                    # Legacy builder is single-line. The exact original is restored
                    # as multiple native PowerPoint paragraphs in patch_pptx_file.
                    obj['text']=obj['text'].replace('\n',' ')
                    obj.pop('rich_runs',None)
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True)
    result=base_builder(prepared,output,mode)
    if mode=='editable':
        extension_report=patch_pptx_file(project,output)
        if isinstance(result,dict):result['runtime_extensions']=extension_report
    return result


def read_native_text(pptx_path:Path):
    """Simple independent shape-text check; formulas must use Studio's OMML QA."""
    from pptx import Presentation
    deck=Presentation(str(pptx_path))
    return [{s.name:s.text for s in slide.shapes if s.has_text_frame} for slide in deck.slides]
