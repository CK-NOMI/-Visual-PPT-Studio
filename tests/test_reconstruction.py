"""v0.2 tests: explicit source binding, geometry, candidate/accept separation, editable output."""
from io import BytesIO
import json
from PIL import Image, ImageDraw
import numpy as np
import pytest
import zipfile
from pptx import Presentation
from fastapi.testclient import TestClient
from pptx import Presentation
from studio import core, server, reconstruction as rc


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(core,'DATA',tmp_path)
    monkeypatch.setattr(server,'DATA',tmp_path)
    return TestClient(server.app)


def make_slide():
    img=Image.new('RGB',(640,360),'white')
    d=ImageDraw.Draw(img)
    d.rounded_rectangle((370,60,560,270),radius=25,fill='#4677DC')
    d.text((70,95),'RESEARCH',fill='black')
    buf=BytesIO();img.save(buf,'PNG');return buf.getvalue()


def make_clean_slide():
    img=Image.new('RGB',(640,360),'white')
    ImageDraw.Draw(img).rounded_rectangle((370,60,560,270),radius=25,fill='#4677DC')
    buf=BytesIO();img.save(buf,'PNG');return buf.getvalue()


def test_ocr_matching_does_not_pair_diagram_digit_with_distant_page_number():
    expected=[{'text':'2','box':[.5275,.1955,.0155,.0287]}]
    observed=[{'text':'2/2','box':[.9402,.9702,.0239,.0276]}]
    match=rc._match_ocr_regions(expected,observed)
    assert len(match)==1
    assert match[0][1] is None


def test_multiscale_ocr_finds_small_text_in_enlarged_slide_tile(monkeypatch):
    image=np.zeros((100,200,3),dtype=np.uint8)
    image[80:95,170:190]=255
    def fake_ocr(value):
        if value.shape[:2]==(108,216) and np.any(value==255):
            return [{'text':'tiny label','box':[.8,.8,.1,.1],'confidence':.91}]
        if value.shape[0]>=200:
            return [{'text':'TITLE','box':[.1,.1,.3,.08],'confidence':.98}]
        return []
    monkeypatch.setattr(rc,'_rapidocr',fake_ocr)
    found=rc._rapidocr_multiscale(image)
    assert {item['text'] for item in found}=={'TITLE','tiny label'}
    tiny=next(item for item in found if item['text']=='tiny label')
    assert tiny['box'][0]>.85
    assert tiny['box'][1]>.85


def make_clean_with_residual_text():
    img=Image.new('RGB',(640,360),'white')
    d=ImageDraw.Draw(img);d.rounded_rectangle((370,60,560,270),radius=25,fill='#4677DC')
    d.text((70,85),'residual',fill='black')
    buf=BytesIO();img.save(buf,'PNG');return buf.getvalue()


def new_project(client):
    p=client.post('/api/projects',json={'title':'Research','topic':'Deep learning','pages':1}).json()
    pid=p['id'];assert client.put(f'/api/projects/{pid}/slides/S001',json={'title':'Research'}).status_code==200
    assert client.post(f'/api/projects/{pid}/lock',json={'locked':True}).status_code==200
    assert client.post(f'/api/projects/{pid}/upload',data={'slide_id':'S001','role':'design'},
       files={'file':('S001_design.png',make_slide(),'image/png')}).status_code==200
    return pid


def test_lab_manual_mask_candidate_accept_and_native_text(client):
    pid=new_project(client);base=f'/api/projects/{pid}/slides/S001'
    region={'text':'RESEARCH','role':'title','box':[.10,.22,.30,.10]}
    res=client.put(base+'/regions',json={'regions':[region]})
    assert res.status_code==200,res.text
    assert client.get(base+'/reconstruction').json()['ocr']['regions'][0]['role']=='title'
    gen=client.post(base+'/clean-candidate',json={'regions':[region],'mask_mode':'rectangle','method':'telea','padding':2})
    assert gen.status_code==200,gen.text
    assert gen.json()['report']['mask_percent']>0
    project=client.get(f'/api/projects/{pid}').json()
    assert 'clean' not in project['slides'][0]['images']  # user approval is required
    assert client.get(f'/lab/{pid}/S001/S001_candidate.png').status_code==200
    assert client.get(f'/lab/{pid}/S001/S002_candidate.png').status_code==404
    approved=client.post(base+'/accept-clean')
    assert approved.status_code==200,approved.text
    export=client.post(f'/api/projects/{pid}/export?mode=editable')
    assert export.status_code==200,export.text
    ppt=core.project_dir(pid)/'build'/'presentation_editable.pptx'
    assert ppt.exists()
    assert any(shape.has_text_frame and 'Research' in shape.text for shape in Presentation(ppt).slides[0].shapes)
    if export.json()['previews']:
        comp=client.post(base+'/compare-render')
        assert comp.status_code==200,comp.text
        assert client.get(f'/lab/{pid}/S001/S001_diff.png').status_code==200


def test_ocr_style_estimation_and_region_persistence(client):
    img=Image.new('RGB',(100,40),'white')
    draw=ImageDraw.Draw(img)
    draw.rectangle((20,10,28,17),fill='#123456')
    estimated=rc._estimate_text_style(np.array(img),[.20,.25,.10,.25])
    assert estimated['color']=='123456'
    assert estimated['size']>0

    pid=new_project(client);base=f'/api/projects/{pid}/slides/S001'
    saved=client.put(base+'/regions',json={'regions':[{
        'text':'Research','box':[.1,.25,.3,.06],'style':{'size':21.5,'color':'123456','bold':True}
    }]})
    assert saved.status_code==200,saved.text
    region=client.get(base+'/reconstruction').json()['ocr']['regions'][0]
    assert region['style']=={'size':21.5,'color':'123456','bold':True}


def test_render_fidelity_score_uses_measured_text_geometry_and_style(monkeypatch,tmp_path):
    image=Image.new('RGB',(120,60),'white')
    draw=ImageDraw.Draw(image);draw.rectangle((20,20,65,30),fill='#123456')
    regions={'source_valid':True,'regions':[{'text':'Research','confidence':.99,
        'box':[20/120,20/60,46/120,11/60], 'style':{'color':'123456'}}]}
    monkeypatch.setattr(rc,'available_ocr',lambda:{'rapidocr':True})
    monkeypatch.setattr(rc,'get_regions',lambda p,sid:regions)
    monkeypatch.setattr(rc,'_rapidocr',lambda arr:[{'text':'Research','confidence':.99,
        'box':[20/120,20/60,46/120,11/60]}])
    result=rc.score_render_fidelity({},'S001',image,image.copy())
    assert result['score']>=95
    assert result['status']=='pass'
    assert set(result['components'])=={'text','position','typography','background','formula','weight'}
    assert result['scoring_version']==4

    regions['regions'][0]['text']='中文标题'
    monkeypatch.setattr(rc,'_rapidocr',lambda arr:[{'text':'', 'confidence':.99,
        'box':[20/120,20/60,46/120,11/60]}])
    cjk_result=rc.score_render_fidelity({},'S001',image,image.copy())
    assert cjk_result['status']=='cjk_text_not_detected'
    assert cjk_result['score']<=49
    assert cjk_result['cjk_expected_regions']==1
    assert cjk_result['cjk_detected_regions']==0

    regions['regions'][0]['text']='P(x)=∑P(y)'
    monkeypatch.setattr(rc,'_rapidocr',lambda arr:[{'text':'P(x)=P(y)','confidence':.99,
        'box':[20/120,20/60,46/120,11/60]}])
    formula_result=rc.score_render_fidelity({},'S001',image,image.copy())
    assert formula_result['formula_verified'] is False
    assert formula_result['score']<95


def test_render_fidelity_penalizes_text_box_width_and_height_mismatch(monkeypatch):
    image=Image.new('RGB',(120,60),'white')
    draw=ImageDraw.Draw(image);draw.rectangle((20,20,65,30),fill='#123456')
    regions={'source_valid':True,'regions':[{'text':'Research','confidence':.99,
        'box':[20/120,20/60,46/120,11/60], 'style':{'color':'123456'}}]}
    monkeypatch.setattr(rc,'available_ocr',lambda:{'rapidocr':True})
    monkeypatch.setattr(rc,'get_regions',lambda p,sid:regions)
    matching=[{'text':'Research','confidence':.99,'box':[20/120,20/60,46/120,11/60]}]
    mismatched=[{'text':'Research','confidence':.99,'box':[20/120,20/60,20/120,5/60]}]
    monkeypatch.setattr(rc,'_rapidocr',lambda arr:matching if arr.shape[:2]==(60,120) else [])
    exact=rc.score_render_fidelity({},'S001',image,image.copy())
    monkeypatch.setattr(rc,'_rapidocr',lambda arr:mismatched if arr.shape[:2]==(60,120) else [])
    off=rc.score_render_fidelity({},'S001',image,image.copy())
    assert off['components']['position']<exact['components']['position']
    assert off['score']<exact['score']

def test_render_fidelity_cannot_pass_with_deterministic_text_overflow(monkeypatch):
    image=Image.new('RGB',(120,60),'white')
    draw=ImageDraw.Draw(image);draw.rectangle((20,20,65,30),fill='#123456')
    region={'text':'研究目标','confidence':.99,'box':[20/120,20/60,46/120,11/60],
            'style':{'color':'123456','size':22,'bold':True}}
    monkeypatch.setattr(rc,'available_ocr',lambda:{'rapidocr':True})
    monkeypatch.setattr(rc,'get_regions',lambda p,sid:{'source_valid':True,'regions':[region]})
    monkeypatch.setattr(rc,'_rapidocr',lambda arr:[{'text':'研究目标','confidence':.99,'box':region['box']}])
    project={'slides':[{'id':'S001','elements':[{'id':'heading','kind':'text','text':'研究目标',
        'box':[.4,.1,.08,.04],'style':{'size':22,'bold':True}}]}]}
    result=rc.score_render_fidelity(project,'S001',image,image.copy())
    assert result['threshold']==90
    assert result['score']<=89
    assert result['layout_warnings'][0]['id']=='heading'
    assert '分数封顶89' in result['notice']


def test_ink_density_detects_bolder_rendered_glyphs(monkeypatch):
    from PIL import ImageFont
    regular=Image.new('RGB',(120,48),'white');bold=Image.new('RGB',(120,48),'white')
    ImageDraw.Draw(regular).text((12,8),'Core',font=ImageFont.load_default(size=28),fill='#17243A')
    ImageDraw.Draw(bold).text((12,8),'Core',font=ImageFont.load_default(size=28),fill='#17243A',stroke_width=1,stroke_fill='#17243A')
    box=[.08,.10,.48,.72]
    assert rc._ink_density(np.asarray(bold),box)>rc._ink_density(np.asarray(regular),box)
    region={'text':'Core','confidence':.99,'box':box,'style':{'color':'17243A'}}
    monkeypatch.setattr(rc,'available_ocr',lambda:{'rapidocr':True})
    monkeypatch.setattr(rc,'get_regions',lambda p,sid:{'source_valid':True,'regions':[region]})
    monkeypatch.setattr(rc,'_rapidocr',lambda arr:[{'text':'Core','confidence':.99,'box':box}])
    scored=rc.score_render_fidelity({},'S001',regular,bold)
    assert any(d['type']=='text_weight' and d['direction']=='actual_heavier'
               for d in scored['differences'])


def test_normalized_stroke_profile_ignores_ocr_box_padding():
    from PIL import ImageFont
    image=Image.new('RGB',(160,54),'white')
    ImageDraw.Draw(image).text((35,13),'Core',font=ImageFont.load_default(size=26),fill='#17243A')
    a=rc._ink_profile(np.asarray(image),[.20,.10,.42,.60])
    b=rc._ink_profile(np.asarray(image),[0,0,1,1])
    assert a['confidence']>=.35 and b['confidence']>=.35
    assert abs(a['stroke_ratio']-b['stroke_ratio'])<.02


def test_uploaded_clean_is_repair_source_and_candidate_is_hash_bound(client):
    pid=new_project(client);base=f'/api/projects/{pid}/slides/S001'
    upload=client.post(f'/api/projects/{pid}/upload',data={'slide_id':'S001','role':'clean'},
        files={'file':('S001_clean.png',make_clean_with_residual_text(),'image/png')})
    assert upload.status_code==200
    region={'text':'residual','box':[.10,.22,.30,.10]}
    result=client.post(base+'/clean-candidate',json={'regions':[region],
        'mask_mode':'rectangle','method':'telea','padding':2,'source_role':'clean'})
    assert result.status_code==200,result.text
    assert result.json()['report']['source_role']=='clean'
    project=client.get(f'/api/projects/{pid}').json()
    assert rc.candidate_is_current(project,'S001',result.json()['report'])
    before=project['slides'][0]['image_metadata']['clean']['sha256']
    accepted=client.post(base+'/accept-clean')
    assert accepted.status_code==200,accepted.text
    after=client.get(f'/api/projects/{pid}').json()['slides'][0]['image_metadata']['clean']['sha256']
    assert after!=before


def test_reconstruction_pipeline_happy_path_renders_and_calibrates(client,monkeypatch):
    """Exercise the complete local path with isolated, deterministic Agent responses."""
    if not server.office_renderer_binary():pytest.skip('没有 LibreOffice，无法验证真实PPTX渲染')
    pid=new_project(client);base=f'/api/projects/{pid}/slides/S001'
    upload=client.post(f'/api/projects/{pid}/upload',data={'slide_id':'S001','role':'clean'},
        files={'file':('S001_clean.png',make_clean_slide(),'image/png')})
    assert upload.status_code==200,upload.text

    async def pair_review(*args,**kwargs):
        return {'score':98,'usable_as_clean':True,'residual_text':[],
                'layout_changes':[],'nontext_damage':[],'recommendations':[],
                'notice':'测试用隔离Agent结果'}
    async def structure(*args,**kwargs):
        return {'objects':[{'id':'title','kind':'text','role':'title','text':'RESEARCH',
                            'box':[.10,.22,.30,.10],'confidence':.99,'style':{}}],
                'notice':'测试结构候选','truncated':False}
    async def visual_review(*args,**kwargs):
        return {'score':96,'issues':[],'suggestions':[]}
    async def calibration(project,slide,design,clean,render,*args):
        assert render.is_file()
        return {'slide_id':'S001','score':90,'issues':['字号层级可微调'],
                'corrections':[{'id':'title','box':[.10,.20,.32,.12],
                                'style':{'size':24,'font':'Microsoft YaHei','italic':False},
                                'reason':'测试校准建议'}],
                'notice':'只提供建议'}
    monkeypatch.setattr(server,'review_image_pair',pair_review)
    monkeypatch.setattr(server,'analyze_structure',structure)
    monkeypatch.setattr(server,'review_visual',visual_review)
    monkeypatch.setattr(server,'calibrate_text_layout',calibration)
    monkeypatch.setattr(rc,'available_ocr',lambda:{'rapidocr':True,'tesseract':False,'simple_lama':False,'opencv':True})
    monkeypatch.setattr(rc,'_rapidocr',lambda arr:[{'text':'RESEARCH','box':[.10,.22,.30,.10],'confidence':.99}])

    reviewed=client.post(base+'/review-image-pair',json={})
    assert reviewed.status_code==200 and reviewed.json()['usable_as_clean']
    assert reviewed.json()['source_pair_current'] is True
    vision=client.post(base+'/analyze-structure',json={})
    assert vision.status_code==200 and vision.json()['objects'][0]['role']=='title'
    ocr=client.post(base+'/ocr?backend=auto')
    assert ocr.status_code==200 and ocr.json()['regions'][0]['confidence']==.99

    region={'text':'RESEARCH','role':'title','editable':True,'confidence':.99,
            'box':[.10,.22,.30,.10]}
    assert client.put(base+'/regions',json={'regions':[region]}).status_code==200
    clean_before=client.get(f'/api/projects/{pid}').json()['slides'][0]['image_metadata']['clean']['sha256']

    element={'id':'title','kind':'text','text':'Research','box':[.10,.22,.30,.10],
             'style':{'size':24,'font':'Microsoft YaHei','color':'17243A'}}
    saved=client.put(base,json={'elements':[element]})
    assert saved.status_code==200,saved.text
    assert client.get(f'/api/projects/{pid}').json()['slides'][0]['image_metadata']['clean']['sha256']==clean_before
    monkeypatch.setenv('PPT_PREVIEW_DISABLE','0')
    exported=client.post(f'/api/projects/{pid}/export?mode=editable')
    assert exported.status_code==200,exported.text
    result=exported.json()
    assert result['native_text']['ok'] is True
    assert result['previews'] and not result['render_warning']
    assert result['visual_comparisons'] and not result['visual_comparisons'][0].get('error')

    compared=client.post(base+'/compare-render')
    assert compared.status_code==200 and compared.json()['diff']=='S001_diff.png'
    assert client.post(base+'/visual-qa',json={}).json()['score']==96
    calibrated=client.post(base+'/calibrate-text',json={})
    assert calibrated.status_code==200 and calibrated.json()['corrections'][0]['style']['font']=='Microsoft YaHei'
    # The API reports corrections but does not silently change the user's layout.
    current=client.get(f'/api/projects/{pid}').json()['slides'][0]['elements'][0]
    assert current['box']==[.10,.22,.30,.10]


def test_manual_asset_extraction_and_bad_geometry(client):
    pid=new_project(client);base=f'/api/projects/{pid}/slides/S001'
    error=client.post(base+'/extract-asset',json={'box':[.9,.2,.2,.2]})
    assert error.status_code==400
    good=client.post(base+'/extract-asset',json={'box':[.58,.18,.29,.5]})
    assert good.status_code==200,good.text
    assert good.json()['filename']=='S001_asset_01.png'
    assert client.get(f'/media/{pid}/S001_asset_01.png').status_code==200
    assert client.post(base+'/clean-candidate',json={'regions':[]}).status_code==400


def test_latex_formula_exports_as_editable_omml(client,monkeypatch):
    if not server.office_renderer_binary():pytest.skip('没有 LibreOffice，无法验证公式的真实页面渲染')
    pid=new_project(client);base=f'/api/projects/{pid}/slides/S001'
    source=r'\frac{QK^T}{\sqrt{d_k}}'
    payload={'id':'attention_formula','kind':'formula','text':source,
             'box':[.2,.4,.6,.2],'style':{'size':26,'align':'center'}}
    saved=client.put(base,json={'elements':[payload]})
    assert saved.status_code==200,saved.text
    preview=client.post('/api/formula/preview',json={'latex':source})
    assert preview.status_code==200 and '<mfrac' in preview.json()['mathml']
    monkeypatch.setenv('PPT_PREVIEW_DISABLE','0')
    export=client.post(f'/api/projects/{pid}/export?mode=editable')
    assert export.status_code==200,export.text
    path=core.project_dir(pid)/'build'/export.json()['filename']
    assert len(Presentation(path).slides)==1
    with zipfile.ZipFile(path) as archive:
        xml=archive.read('ppt/slides/slide1.xml').decode('utf-8')
    assert 'a14:m' in xml and 'm:oMathPara' in xml and ':f>' in xml, (xml.count('a14:m'),xml.count('oMathPara'),xml.count(':f>'))
    assert export.json()['native_text']['slides'][0]['native_equation_count']==1
    assert export.json()['native_text']['slides'][0]['equations_pass'] is True
    assert export.json()['previews'] and not export.json()['render_warning']
    rendered=Image.open(core.project_dir(pid)/'build'/export.json()['previews'][0]).convert('RGB')
    crop=rendered.crop((int(rendered.width*.2),int(rendered.height*.4),int(rendered.width*.8),int(rendered.height*.6)))
    pixels=list(crop.getdata());base_color=pixels[0]
    ink_pixels=sum(1 for px in pixels if max(abs(px[i]-base_color[i]) for i in range(3))>45)
    assert ink_pixels>50,'LibreOffice 真实渲染中没有看到公式笔画'


def test_formula_syntax_and_length_are_validated(client):
    pid=new_project(client);base=f'/api/projects/{pid}/slides/S001'
    bad={'id':'bad_formula','kind':'formula','text':r'\frac{a}','box':[.1,.1,.8,.2]}
    assert client.put(base,json={'elements':[bad]}).status_code==400
    assert client.post('/api/formula/preview',json={'latex':'x'*2001}).status_code==422


def test_clean_candidate_invalidated_if_source_changes(client):
    pid=new_project(client);base=f'/api/projects/{pid}/slides/S001'
    assert client.post(base+'/clean-candidate',json={'regions':[{'box':[.1,.2,.3,.1]}], 'mask_mode':'rectangle'}).status_code==200
    new=Image.new('RGB',(640,360),(240,240,240));buf=BytesIO();new.save(buf,'PNG')
    assert client.post(f'/api/projects/{pid}/upload',data={'slide_id':'S001','role':'design'},
        files={'file':('another.png',buf.getvalue(),'image/png')}).status_code==200
    assert client.post(base+'/accept-clean').status_code==400
    assert client.get(base+'/reconstruction').json()['ocr'] is None


def test_ocr_engine_mock_and_region_bounds(client,monkeypatch):
    pid=new_project(client);base=f'/api/projects/{pid}/slides/S001'
    monkeypatch.setattr(rc,'available_ocr',lambda:{'rapidocr':True,'tesseract':False,'simple_lama':False,'opencv':True})
    monkeypatch.setattr(rc,'_rapidocr',lambda arr:[{'text':'RESEARCH','box':[.08,.25,.20,.05],'confidence':.96}])
    o=client.post(base+'/ocr?backend=rapidocr');assert o.status_code==200,o.text
    assert o.json()['regions'][0]['text']=='RESEARCH'
    assert client.put(base+'/regions',json={'regions':[{'text':'X','box':[-.01,0,.4,.2]}]}).status_code==400
    assert client.put(base+'/regions',json={'regions':[{'text':'Research','box':[.1,.25,.3,.06],'role':'title'}]}).status_code==200
    assert client.get(base+'/reconstruction').json()['ocr']['source_valid']
