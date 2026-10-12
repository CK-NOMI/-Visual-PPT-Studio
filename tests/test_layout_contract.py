from copy import deepcopy
from io import BytesIO
from pathlib import Path
import json
import zipfile
import pytest
from PIL import Image
from pptx import Presentation
from fastapi.testclient import TestClient
from studio import core,server,reconstruction as rc
from studio.text_layout import layout_project
from studio.pptx_native import build
from studio.quality import check_native_text,check_saved_layout


def project():
    return {'id':'test','revision':1,'design':{'bg':'FFFFFF','ink':'17243A','font':'Noto Sans'},
            'slides':[{'id':'S001','background':'FFFFFF','images':{},'elements':[
                {'id':'mixed','kind':'text','text':'  中文 English 中文  ','box':[.1,.1,.7,.08],
                 'style':{'size':20,'font':'Noto Sans','eastAsianFont':'Noto Sans SC','margin_left_pt':1.5}}
            ]}]}


def test_pptx_is_exact_native_single_line_and_contains_layout_manifest(tmp_path):
    p,_=layout_project(project());before=deepcopy(p);file=tmp_path/'layout.pptx';build(p,file)
    assert p==before
    assert check_native_text(p,file)['ok']
    assert check_saved_layout(p,file)['ok']
    pp=Presentation(file)
    assert len(pp.slides[0].shapes)==len(p['slides'][0]['elements'])
    assert ''.join(sh.text for sh in pp.slides[0].shapes)=='  中文 English 中文  '
    assert all('\n' not in sh.text and sh.text_frame.word_wrap is False for sh in pp.slides[0].shapes)
    with zipfile.ZipFile(file) as archive:
        xml=archive.read('customXml/visualppt-layout.xml').decode()
        assert 'stage_visual_renderer' in xml and 'file_sha256' in xml


@pytest.mark.parametrize('mutation',['space','size','duplicate','box'])
def test_exact_check_detects_corruption_not_hidden_by_loose_text_matching(tmp_path,mutation):
    from copy import deepcopy as dc
    from pptx.util import Pt
    p,_=layout_project(project());file=tmp_path/'layout.pptx';build(p,file);pp=Presentation(file)
    shape=pp.slides[0].shapes[0]
    if mutation=='space':shape.text=shape.text.strip()
    elif mutation=='size':shape.text_frame.paragraphs[0].runs[0].font.size=Pt(10)
    elif mutation=='box':shape.left+=100000
    else:pp.slides[0].shapes._spTree.append(dc(shape._element))
    pp.save(file)
    assert not check_saved_layout(p,file)['ok']


def test_new_layout_route_is_pure_and_saving_reloading_preserves_it(tmp_path,monkeypatch):
    monkeypatch.setattr(core,'DATA',tmp_path);monkeypatch.setattr(server,'DATA',tmp_path)
    p=core.create_project('Layout','Mixed',2);p['locked']=True;core.save_project(p)
    with TestClient(server.app) as client:
        path=f'/api/projects/{p["id"]}/slides/S001'
        proposal=client.post(path+'/text-layout',json={'elements':project()['slides'][0]['elements']})
        assert proposal.status_code==200,proposal.text
        els=proposal.json()['elements']
        assert core.read_project(p['id'])['revision']==p['revision']
        response=client.put(path,json={'elements':els});assert response.status_code==200,response.text
        assert response.json()['elements']==els
        current=client.get(f'/api/projects/{p["id"]}').json()
        assert current['slides'][0]['elements']==els
        assert current['slides'][1]['elements']==p['slides'][1]['elements']
        font=els[0]['text_layout']['font']
        result=client.get(f'/api/projects/{p["id"]}/fonts/{font["file_identity"]}.ttf')
        assert result.status_code==200 and len(result.content)>10000
        assert client.get(f'/api/projects/{p["id"]}/fonts/not-a-registered-face.ttf').status_code==404


def test_more_than_old_148_objects_are_not_silently_dropped(tmp_path,monkeypatch):
    monkeypatch.setattr(core,'DATA',tmp_path)
    p=core.create_project('Labels','Labels',1)
    els=[{'id':f'label{i}','kind':'text','text':'Token1' if i%2 else 'Token2',
          'box':[.01+(i%15)*.063,.01+(i//15)*.07,.06,.035],'style':{'size':8,'font':'Arial'}} for i in range(180)]
    result=core.update_slide(p,'S001',{'elements':els})
    assert len(result['elements'])==180
    assert len({e['run_id'] for e in result['elements']})==180
