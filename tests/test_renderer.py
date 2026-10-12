from io import BytesIO
import zipfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from studio import core,server


def test_office_renderer_binary_accepts_libreoffice_and_soffice(monkeypatch):
    bins={'libreoffice':None,'soffice':'/opt/test/bin/soffice'}
    monkeypatch.setattr(server.shutil,'which',lambda name:bins.get(name))
    assert server.office_renderer_binary()=='/opt/test/bin/soffice'
    bins['libreoffice']='/opt/test/bin/libreoffice'
    assert server.office_renderer_binary()=='/opt/test/bin/libreoffice'


def asset():
    im=Image.new('RGBA',(300,200),(255,100,25,210));b=BytesIO();im.save(b,'PNG');return b.getvalue()

def test_native_shapes_and_raster_fallback(tmp_path,monkeypatch):
    monkeypatch.setattr(core,'DATA',tmp_path)
    monkeypatch.setattr(server,'DATA',tmp_path)
    p=core.create_project('Test','Topic',1,'dark')
    p['locked']=True;core.save_project(p)
    core.write_image(p,'S001','design',asset(),'S001_design.png')
    core.write_image(p,'S001','asset_01',asset(),'S001_asset_01.png')
    s=p['slides'][0]
    new=s['elements']+[
      {'id':'square','kind':'rect','box':[.75,.2,.14,.21],'text':'','style':{'fill':'4455AA'}},
      {'id':'circle','kind':'ellipse','box':[.7,.5,.15,.2],'text':'','style':{'fill':'77AAFF'}},
      {'id':'separator','kind':'line','box':[.07,.7,.35,.01],'text':'','style':{'stroke':'FFFFFF','strokeWidth':2}},
      {'id':'picture','kind':'image','box':[.53,.24,.17,.2],'asset':'S001_asset_01.png','style':{}}
    ]
    core.update_slide(p,'S001',{'elements':new})
    er=server.render_deck(p['id'],'editable')
    rr=server.render_deck(p['id'],'raster')
    assert er['native_text']['ok']
    eb=core.project_dir(p['id'])/'build'/er['filename']
    rb=core.project_dir(p['id'])/'build'/rr['filename']
    assert len(Presentation(eb).slides[0].shapes)>=6
    with zipfile.ZipFile(eb) as z:
        media=[n for n in z.namelist() if n.startswith('ppt/media/') and not n.endswith('/')]
        # Only the independent asset image is included, not design.png.
        assert len(media)==1
    with zipfile.ZipFile(rb) as z:
        media=[n for n in z.namelist() if n.startswith('ppt/media/') and not n.endswith('/')]
        assert len(media)==1
