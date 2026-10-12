from pptx import Presentation
from pptx.util import Inches

from studio.fonts import font_choices
from studio.pptx_native import _set_run_fonts
from studio.pptx_native import build
from studio.server import office_font_environment, office_renderer_binary
import tempfile


def test_text_runs_write_explicit_latin_and_east_asian_font_slots():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = '中文标题 Fast Inference'

    _set_run_fonts(run, 'Arial', 'SimSun')

    xml = run._r.xml
    assert 'typeface="Arial"' in xml
    assert 'typeface="SimSun"' in xml
    assert '<a:ea' in xml
    assert '<a:cs' in xml


def test_font_choices_are_named_common_families():
    choices = font_choices()
    assert choices['latin']
    assert choices['east_asian']
    assert 'Noto Sans SC' in choices['east_asian']
    assert 'Noto Serif SC' in choices['east_asian']
    assert 'Noto Sans' in choices['latin']
    assert 'Noto Serif' in choices['latin']
    assert all(isinstance(name, str) and name for group in choices.values() for name in group)


def test_packaged_variable_font_families_exist():
    from pathlib import Path
    font_dir = Path(__file__).parents[1] / 'studio' / 'assets' / 'fonts'
    for filename in ('NotoSans-VF.ttf', 'NotoSerif-VF.ttf', 'NotoSansSC-VF.ttf', 'NotoSerifSC-VF.ttf'):
        assert (font_dir / filename).is_file()
    for filename in ('OFL-Noto-Sans.txt', 'OFL-Noto-Serif.txt',
                     'OFL-Noto-Sans-SC.txt', 'OFL-Noto-Serif-SC.txt'):
        assert 'SIL OPEN FONT LICENSE' in (font_dir / filename).read_text(encoding='utf-8')


def test_office_renderer_is_told_where_packaged_fonts_live():
    with tempfile.TemporaryDirectory() as tmp:
        env = office_font_environment(office_renderer_binary() or 'soffice', tmp)
        assert env['FONTCONFIG_FILE']
        with open(env['FONTCONFIG_FILE'], encoding='utf-8') as config:
            assert 'studio/assets/fonts' in config.read()


def test_native_text_exports_tracking_and_line_spacing(tmp_path):
    output = tmp_path / 'spacing.pptx'
    project = {
        'design': {'bg': 'FFFFFF', 'ink': '17243A', 'font': 'Arial'},
        'slides': [{'id': 'S001', 'background': 'FFFFFF', 'images': {}, 'elements': [
            {'id': 't1', 'kind': 'text', 'text': '中文 English', 'box': [.1, .1, .8, .3],
             'style': {'size': 18, 'char_spacing_pt': 1.25, 'line_spacing': 1.3}}
        ]}]
    }
    build(project, output)
    presentation = Presentation(output)
    paragraph = presentation.slides[0].shapes[0].text_frame.paragraphs[0]
    assert paragraph.line_spacing == 1.3
    assert paragraph.runs[0]._r.get_or_add_rPr().get('spc') == '125'
