from pathlib import Path
import hashlib
from pptx import Presentation
import pytest
from fastapi.testclient import TestClient

from studio.pptx_native import build
from studio.quality import check_native_text, check_saved_layout
from studio.visual_ppt_runtime.export_bridge import build_with_extensions
from studio.visual_ppt_runtime.tools import EditingTools
from studio.visual_ppt_runtime.transactions import EditingSession
from studio.visual_ppt_runtime.permissions import PermissionPolicy


def test_runtime_export_bridge_preserves_precise_rich_text_and_native_omml(tmp_path):
    project={'id':'82624a838e29','revision':1,'design':{'bg':'FFFFFF','ink':'17243A','font':'Arial'},
        'slides':[{'id':'S001','background':'FFFFFF','images':{},'elements':[
            {'id':'mixed','kind':'text','text':'标题 Fast Inference','box':[.1,.1,.8,.1],
             'style':{'font':'Arial','eastAsianFont':'Noto Sans SC','size':22,
                      'char_spacing_pt':.35,'line_spacing':1.2},
             'rich_runs':[{'text':'标题 ','style':{'font':'Arial','eastAsianFont':'Noto Sans SC'}},
                          {'text':'Fast Inference','style':{'font':'Arial','size':22,'bold':True,
                              'char_spacing_pt':.35}}]},
            {'id':'math','kind':'formula','text':'x^2+1','box':[.1,.3,.3,.1],
             'style':{'size':22,'align':'center'}}]}]}
    target=tmp_path/'runtime-export.pptx'
    build_with_extensions(project,target,base_builder=build)
    assert check_native_text(project,target)['ok']
    assert check_saved_layout(project,target)['ok']
    deck=Presentation(target);runs=deck.slides[0].shapes[0].text_frame.paragraphs[0].runs
    assert [r.text for r in runs]==['标题 ','Fast Inference']
    assert runs[1].font.size.pt==22 and runs[1].font.bold is True
    assert runs[1]._r.get_or_add_rPr().get('spc')=='35'


def test_runtime_page_state_bounds_model_pagination_and_reads_scene():
    project={'id':'82624a838e29','revision':1,'design':{'bg':'FFFFFF','ink':'17243A','font':'Arial'},
        'slides':[{'id':'S001','background':'FFFFFF','images':{},'elements':[
            {'id':f't{i}','kind':'text','text':str(i),'box':[.1,.1,.1,.1],'style':{'size':18}}
            for i in range(3)]}]}
    session=EditingSession(project,PermissionPolicy(mode='free',allowed_slides=frozenset({'S001'})))
    result=EditingTools(session).dispatch('get_page_state',{'slide_id':'S001','offset':'bad','limit':1000})
    assert result['ok'] is True
    assert result['result']['total']==3
    assert len(result['result']['elements'])==3


def test_editing_runtime_http_prepare_isolated_candidate_without_writing_project(tmp_path,monkeypatch):
    from io import BytesIO
    from PIL import Image
    from studio import core,server
    monkeypatch.setattr(core,'DATA',tmp_path)
    monkeypatch.setattr(server,'DATA',tmp_path)
    if not server.office_renderer_binary():pytest.skip('LibreOffice required for runtime HTTP render check')
    project=core.create_project('Runtime HTTP Integration','isolated runtime candidate',pages=1)
    pid=project['id'];page=project['slides'][0]
    page['images']['design']='S001_design.png'
    image=BytesIO();Image.new('RGB',(640,360),'white').save(image,'PNG')
    (core.project_dir(pid)/'images'/'S001_design.png').write_bytes(image.getvalue())
    project['locked']=True;core.save_project(project)
    source=core.project_dir(pid)/'project.json';before=hashlib.sha256(source.read_bytes()).hexdigest()
    client=TestClient(server.app)
    created=client.post(f'/api/projects/{pid}/editing-runtime/sessions',json={
        'slide_ids':['S001'],'mode':'free','allow_text_changes':False})
    assert created.status_code==200,created.text
    session_id=created.json()['session_id']
    from studio.editing_runtime_api import SESSIONS
    session=SESSIONS[f'{pid}:{session_id}']['session']
    tx=session.execute('S001','font-size-check',[{'action':'set_properties','element_id':'title',
        'changes':{'style':{'size':26,'char_spacing_pt':.35,'line_spacing':1.05}}}])
    decision=client.post(f'/api/projects/{pid}/editing-runtime/sessions/{session_id}/transactions/{tx["tx_id"]}/decision',
        json={'accepted':True})
    assert decision.status_code==200,decision.text
    prepared=client.post(f'/api/projects/{pid}/editing-runtime/sessions/{session_id}/prepare',json={})
    assert prepared.status_code==200,prepared.text
    report=prepared.json()
    assert report['status']=='prepared' and report['renderer']=='LibreOffice'
    assert report['native_text']['ok'] and report['saved_layout_check']['ok']
    assert (tmp_path/'editing_runtime'/pid/session_id/'prepared'/'candidate.pptx').is_file()
    assert hashlib.sha256(source.read_bytes()).hexdigest()==before
