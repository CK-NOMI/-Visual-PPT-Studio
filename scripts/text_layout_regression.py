"""Local, reproducible before/after artifacts; never edits the source workspace.

Usage: .venv/bin/python scripts/text_layout_regression.py PROJECT_ID --label baseline
"""
from __future__ import annotations

import argparse
from copy import deepcopy
import json
from pathlib import Path
import subprocess
import sys
import tempfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from studio.core import read_project, project_dir, safe_json_write
from studio.pptx_native import build
from studio.server import office_renderer_binary, office_font_environment
from studio.quality import check_native_text


def render_files(files, folder):
    lo = office_renderer_binary()
    result = {'renderer': 'LibreOffice auxiliary preview', 'office_verified': False,
              'powerpoint_installed': Path('/Applications/Microsoft PowerPoint.app').exists(),
              'previews': [], 'errors': []}
    if not lo:
        result['errors'].append('LibreOffice unavailable')
        return result
    import pymupdf
    with tempfile.TemporaryDirectory(prefix='layout-preview-') as tmp:
        proc = subprocess.run([lo, '-env:UserInstallation=file://' + tmp + '/profile',
                               '--headless', '--convert-to', 'pdf', '--outdir', str(folder),
                               *map(str, files)], capture_output=True, text=True, timeout=120,
                              env=office_font_environment(lo, tmp))
        result['conversion_output'] = (proc.stdout + proc.stderr)[-1500:]
    for file in files:
        pdf = file.with_suffix('.pdf')
        if not pdf.exists():
            result['errors'].append('Missing PDF: ' + file.name)
            continue
        with pymupdf.open(pdf) as doc:
            for i, page in enumerate(doc, 1):
                path = folder / f'{file.stem}_S{i:03d}.png'
                page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False).save(path)
                result['previews'].append(path.name)
    return result


def trace_shapes(project, file):
    from pptx import Presentation
    from pptx.oxml.ns import qn
    pp = Presentation(file)
    pages = []
    for source, slide in zip(project['slides'], pp.slides):
        objects = []
        expected = {e['id']: e for e in source['elements']}
        for z, shape in enumerate(slide.shapes):
            if not shape.has_text_frame:
                continue
            original = expected.get(shape.name, {})
            runs = []
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    props = r._r.get_or_add_rPr()
                    runs.append({'text': r.text, 'size_pt': r.font.size.pt if r.font.size else None,
                                 'font': r.font.name, 'eastAsianFont': props.find(qn('a:ea')).get('typeface')
                                 if props.find(qn('a:ea')) is not None else None})
            objects.append({'id': shape.name, 'shape_id': shape.shape_id, 'z_order': z,
                            'source': original,
                            'pptx': {'text': shape.text, 'box': [shape.left/pp.slide_width,
                                      shape.top/pp.slide_height, shape.width/pp.slide_width,
                                      shape.height/pp.slide_height], 'runs': runs,
                                     'word_wrap': shape.text_frame.word_wrap}})
        pages.append({'slide_id': source['id'], 'objects': objects})
    return pages


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('project_id')
    parser.add_argument('--label', default='candidate')
    parser.add_argument('--no-render', action='store_true')
    args = parser.parse_args()
    if not args.label.replace('-', '').isalnum():
        parser.error('label must be alphanumeric')
    original = read_project(args.project_id)
    folder = project_dir(args.project_id)/'build'/'text-layout-v2'/args.label
    folder.mkdir(parents=True, exist_ok=True)
    if (folder/'full.pptx').exists():
        parser.error('This artifact label already exists; use a new --label to preserve its evidence/hash')
    safe_json_write(folder/'source_project.json', original)
    project = deepcopy(original)
    if args.label != 'baseline':
        from studio.text_layout import layout_project
        project, diagnostics = layout_project(project)
        safe_json_write(folder/'layout_report.json', diagnostics)
    safe_json_write(folder/'layout_project.json', project)
    files = []
    for mode in ('full', 'clean', 'native'):
        variant = deepcopy(project)
        for slide in variant['slides']:
            if mode == 'clean':
                slide['elements'] = []
            if mode == 'native':
                slide['images'] = {}
                slide['background'] = 'FFFFFF'
                slide['elements'] = [e for e in slide['elements'] if e['kind'] in ('text', 'formula')]
        file = folder/f'{mode}.pptx'
        build(variant, file)
        files.append(file)
    safe_json_write(folder/'shape_trace.json', trace_shapes(project, files[0]))
    safe_json_write(folder/'native_check.json', check_native_text(project, files[0]))
    if not args.no_render:
        safe_json_write(folder/'render_report.json', render_files(files, folder))
    print(folder)


if __name__ == '__main__':
    main()
